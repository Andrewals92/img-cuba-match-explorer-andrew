import openpyxl, json, re, collections, datetime
from pathlib import Path
import argparse
ap=argparse.ArgumentParser();ap.add_argument('--directory-json',required=True);ap.add_argument('--residency-explorer',required=True);ap.add_argument('--match-a-resident',required=True);ap.add_argument('--output',required=True);args=ap.parse_args()
ROOT=Path(__file__).parent
D=json.loads((Path(args.directory_json)).read_text())
def norm(s): return re.sub(r'[^a-z0-9]', '', str(s).lower().replace('program',''))
idx=collections.defaultdict(list)
for p in D: idx[(norm(p['name']),norm(p['city']))].append(p)
out=[]
w=openpyxl.load_workbook(Path(args.residency_explorer),read_only=True,data_only=True)
r=list(w['Programas'].iter_rows(values_only=True));h=r[5]
for row in r[6:]:
 if not row[0]: continue
 city,state=str(row[1]).rsplit(',',1);matches=idx[(norm(row[0]),norm(city))]
 out.append({'source_key':row[18],'name':row[0],'city':city.strip(),'state':state.strip(),'specialty':'Internal Medicine','acgme_program_id':matches[0]['acgme_program_id'] if len(matches)==1 else None,'source':'Residency Explorer','source_cycle':2026,'source_date':None,'program_url':row[18],'fields':{h[i]:row[i] for i in [2,7,8,9,10,11,12,13,14,15,16,17] if row[i] is not None}})
w=openpyxl.load_workbook(Path(args.match_a_resident),read_only=True,data_only=True)
r=list(w['Programas'].iter_rows(values_only=True));h=r[4]
for row in r[5:]:
 if not row[1]: continue
 out.append({'source_key':'MAR:'+re.sub(r'\D','',str(row[1])),'name':row[2],'state':row[3],'specialty':'Internal Medicine','acgme_program_id':re.sub(r'\D','',str(row[1])),'source':'Match A Resident','source_date':'2026-09-11','source_cycle':None,'program_url':'https://www.matcharesident.com/specialty-navigator/internal-medicine/','fields':{h[i]:row[i] for i in [6,7,8,9,10,11,12,13,14,15,16,17,18,19,20,21,29] if row[i] is not None}})
# Only program-owned detail sections; connections/surveys/comments are excluded.
extra=collections.defaultdict(dict)
for row in list(w['Detalles'].iter_rows(values_only=True))[5:]:
 if row[3] in ['CORE COMPETENCIES','SUPPLEMENTAL INFO','PROGRAM HIGHLIGHTS'] and row[5] is not None:
  extra[re.sub(r'\D','',str(row[1]))][str(row[3])+' · '+str(row[4])]=str(row[5])
for entry in out:
 if entry['source']=='Match A Resident': entry['fields'].update(extra[entry['acgme_program_id']])
# Program-only application trends; these are reference-sample shares, not interview probabilities.
for row in r[5:]:
 if not row[1]: continue
 entry=next(x for x in out if x['source_key']=='MAR:'+re.sub(r'\D','',str(row[1])))
 for i,label in [(26,'Tendencia de aplicaciones · temporada previa (muestra de referencia)'),(27,'Tendencia de aplicaciones · temporada actual (muestra de referencia)')]:
  if row[i] is not None:entry['fields'][label]=row[i]
for row in list(w['Detalles'].iter_rows(values_only=True))[5:]:
 if row[3]=='ENLACES' and row[4]=='Watch Program Overview Video' and str(row[5]).startswith('https://'):
  entry=next(x for x in out if x['source_key']=='MAR:'+re.sub(r'\D','',str(row[1])))
  entry['fields']['Vídeo de presentación del programa']=row[5]
# Apply the audited 19 mappings by immutable source key; no generated program UUIDs.
resolved=ROOT/'program-guide-v43'/'identity-resolutions.json'
if resolved.exists():
 mapping={x['source_key']:x['acgme_program_id'] for x in json.loads(resolved.read_text())}
 for entry in out:
  if entry['source_key'] in mapping:entry['acgme_program_id']=mapping[entry['source_key']]
def encode(v):
 if isinstance(v,(datetime.date,datetime.datetime)): return v.isoformat()
 raise TypeError(type(v).__name__)
(Path(args.output)).write_text(json.dumps(out,ensure_ascii=False,indent=2,default=encode))
print({'RE_records':sum(p['source']=='Residency Explorer' for p in out),'RE_identity_matches':sum(p['source']=='Residency Explorer' and bool(p['acgme_program_id']) for p in out),'MAR_records':sum(p['source']=='Match A Resident' for p in out)})
